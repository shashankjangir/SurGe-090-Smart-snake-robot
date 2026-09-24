"""Slide 16 - 'What Changed', duplicated from slide 1 (the only dark slide).

Slide 1 is a title slide, so this is the one slide that needs real
reconstruction rather than text substitution:
  keep    the two decorative circles, the chip, the title block, the logo
  reuse   the three team cards as the three timeline stations
  drop    institution / guide / repo / centre icon
  add     the corrections register table (the reason this slide exists)
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from edit import sl, find, txt, kill, notes

P = "work.pptx"
prs = Presentation(P)
s = sl(prs, 16)

INK      = "E2E8F0"   # body text on dark
MUTED    = "6B7890"
TEAL     = "0E96AC"
CYAN     = "2BD4E8"
CARDFILL = "1A2440"
EDGE     = "2A3958"
AMBER    = "E8821E"

# --- strip the title-slide furniture that has no role here -----------------
kill(s, "Text 4", "Shape 14", "Image 0",
     "Text 30", "Text 31", "Text 32", "Text 33",
     "RepoLabel", "RepoURL")

# --- header ----------------------------------------------------------------
txt(find(s, "Text 3"), "WHAT CHANGED SINCE THIS REVIEW")

title = find(s, "Text 5")
title.top, title.height = Inches(1.28), Inches(0.86)
txt(title, "Generation 2 — Corrections & Cause")

sub = find(s, "Text 6")
sub.top, sub.height, sub.width = Inches(2.12), Inches(0.42), Inches(11.8)
txt(sub, "Slides 1–15 are unchanged from the 25% review. This slide records what has since been superseded.")

# --- corrections register --------------------------------------------------
ROWS = [
    ("SLIDE", "THAT DECK SAID", "TODAY", "WHY"),
    ("10", "XL330 · U2D2 · Power Hub — “ORDERED, AWAITING DELIVERY”",
     "Order cancelled. 10× ST3215 + Waveshare ESP32 driver in hand",
     "Part unobtainable in India"),
    ("12", "“Eight Phases · Currently at 25%”",
     "Phase 3 — Fabrication & Assembly",
     "Eight weeks of Gen 2 work since"),
    ("13", "“Prototype fabrication started”",
     "Gen 1 segments printed; Gen 2 v6 not yet printed",
     "Design of record changed"),
    ("15", "Roadmap step 1 — “Receive Dynamixel XL330 motors”",
     "Superseded — see slide 29",
     "Different motors received"),
]
COLW = [0.85, 4.55, 4.35, 2.42]          # sums to 12.17
LEFT, TOP, HEIGHT = Inches(0.58), Inches(2.66), Inches(2.46)

gf = s.shapes.add_table(len(ROWS), 4, LEFT, TOP, Inches(12.17), HEIGHT)
gf.name = "CorrectionsRegister"
tbl = gf.table
tbl.first_row = False           # suppress the built-in light banding
tbl.horz_banding = False
for i, w in enumerate(COLW):
    tbl.columns[i].width = Inches(w)

for ri, row in enumerate(ROWS):
    tbl.rows[ri].height = Inches(0.62 if ri == 0 else 0.46)
    for ci, val in enumerate(row):
        c = tbl.cell(ri, ci)
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor.from_string(CARDFILL if ri == 0 else "16223A")
        c.margin_left, c.margin_right = Inches(0.10), Inches(0.08)
        c.margin_top, c.margin_bottom = Inches(0.04), Inches(0.04)
        tf = c.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = val
        f = r.font
        f.name = "Calibri"
        if ri == 0:
            f.size, f.bold = Pt(10), True
            f.color.rgb = RGBColor.from_string(TEAL)
        elif ci == 0:
            f.size, f.bold = Pt(14), True
            f.color.rgb = RGBColor.from_string(CYAN)
        else:
            f.size = Pt(9.5)
            f.bold = (ci == 2)
            f.color.rgb = RGBColor.from_string(INK if ci == 2 else MUTED)

# --- three timeline stations (reusing the team cards) ----------------------
STATIONS = [
    ("Text 17", "Text 18", "Text 19",
     "Constraint", "WHAT HAPPENED", "2 of 10 XL330 received, then the part vanished"),
    ("Text 22", "Text 23", "Text 24",
     "Six-week stall", "IMPACT", "Control stack written against absent hardware"),
    ("Text 27", "Text 28", "Text 29",
     "Re-architecture", "COMMIT a32b675 · 2026-08-20", "“No dynamixel in the market. So switched to Servos”"),
]
for head_n, lab_n, det_n, head, lab, det in STATIONS:
    txt(find(s, head_n), head)
    txt(find(s, lab_n), lab)
    txt(find(s, det_n), det)

# --- closing callout -------------------------------------------------------
box = s.shapes.add_textbox(Inches(0.85), Inches(6.94), Inches(11.6), Inches(0.42))
box.name = "PivotCallout"
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
for text, colour, bold in (
        ("WHAT IT BOUGHT   ", AMBER, True),
        ("5.7× joint torque (0.52 → 2.94 N·m)  ·  a 12 V rail  ·  sourceable in India  —  "
         "consequences of the switch, not its reason. Torque budget: slide 23.", INK, False)):
    r = p.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold = "Calibri", Pt(10.5), bold
    r.font.color.rgb = RGBColor.from_string(colour)

notes(s, "Open by saying the deck up to this point is unchanged from the 25 "
         "percent review and is being left that way deliberately - those slides "
         "are a record of what we believed then, not a claim about now. Then "
         "walk the register top to bottom. It takes about forty seconds and it "
         "pre-empts every 'but slide 10 says' interruption in the second half. "
         "Say the honest version of the pivot out loud: this was a procurement "
         "failure, not an engineering preference, and a re-architecture under a "
         "supply constraint is a legitimate result. The follow-up question a "
         "reviewer asks is 'where is your torque budget' - it is slide 23, and "
         "it did not exist at the 25 percent review.")

prs.save(P)
print("slide 16 written")
