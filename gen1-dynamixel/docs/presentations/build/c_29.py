"""Slide 29 - roadmap, duplicated from slide 12 (8-phase alternating timeline)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from edit import sl, find, txt, fill, notes

P = "work.pptx"
prs = Presentation(P)
s = sl(prs, 29)

BLUE, AMBER = "1E6FD6", "E8821E"

txt(find(s, "Text 0"), "ROADMAP")
txt(find(s, "Text 1"), "Path to Demonstration — Generation 2")

# phase label / name / status shape names, in timeline order
PHASES = [
    ("Text 12", "Text 13", "Text 14", "Shape 15",
     "PHASE 1", "Power Verification", "In progress", BLUE),
    ("Text 19", "Text 20", "Text 21", "Shape 22",
     "PHASE 2", "Servo Bring-Up", "Up next", BLUE),
    ("Text 26", "Text 27", "Text 28", "Shape 29",
     "PHASE 3", "Segment Fabrication", "Planned", AMBER),
    ("Text 33", "Text 34", None, "Shape 35",
     "PHASE 4", "Mechanical Assembly", None, AMBER),
    ("Text 39", "Text 40", None, "Shape 41",
     "PHASE 5", "Field Integration", None, AMBER),
    ("Text 45", "Text 46", None, "Shape 47",
     "PHASE 6", "Locomotion Trials", None, AMBER),
    ("Text 51", "Text 52", None, "Shape 53",
     "PHASE 7", "Characterisation", None, AMBER),
    ("Text 57", "Text 58", None, "Shape 59",
     "PHASE 8", "Final Demonstration", None, AMBER),
]

for lab_n, name_n, stat_n, circ_n, lab, name, stat, colour in PHASES:
    txt(find(s, lab_n), lab)
    txt(find(s, name_n), name)
    if stat_n and stat:
        txt(find(s, stat_n), stat)
    fill(find(s, circ_n), colour)

# "WE ARE HERE" sits under phase 2 in the source; phase 1 is current here.
for nm in ("Shape 61", "Text 62"):
    find(s, nm).left = Inches(0.55)
txt(find(s, "Text 62"), "WE ARE HERE")

# The freeze leaves the obsolete Gen 1 roadmap standing on slide 15, so this
# slide has to say out loud that it supersedes it.
box = s.shapes.add_textbox(Inches(0.58), Inches(1.44), Inches(12.17), Inches(0.30))
box.name = "SupersedeNote"
tf = box.text_frame
tf.word_wrap = True
r = tf.paragraphs[0].add_run()
r.text = ("Replaces the Generation 1 roadmap on slide 15, which was written "
          "around Dynamixel XL330 delivery and is now obsolete.")
r.font.size = Pt(10)
r.font.name = "Calibri"
r.font.italic = True
r.font.color.rgb = RGBColor.from_string("6B7890")

notes(s, "Open by naming slide 15 and saying plainly that it is superseded - do "
         "not let the audience find two roadmaps on their own and wonder which "
         "one is live. The eight phases here are the actual next steps from the "
         "hardware inventory, not a re-skin of the old plan. Phase 1 is where we "
         "are: the pack cannot go on the robot until the fuse, loop key and BMS "
         "are fitted. Close on phase 3, because it contains the one experiment "
         "everything else depends on - printing a v6 segment and measuring "
         "whether the belly friction is actually anisotropic. If it is not, the "
         "gait does not work and we would rather know that before printing ten "
         "more segments. Target for the next review is first translating "
         "locomotion on hardware, which is a milestone you can watch rather "
         "than a percentage you have to trust.")

prs.save(P)
print("slide 29 written")
