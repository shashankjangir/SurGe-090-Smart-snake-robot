"""Helpers for rewriting text in duplicated slides without losing formatting.

python-pptx's ``text_frame.text = ...`` collapses every run to one unstyled run,
which destroys the deck's type scale. Everything here assigns ``run.text``
instead and clones the source paragraph's <a:pPr> so spacing and bullets carry.
"""
import copy


def sl(prs, n):
    """1-indexed slide access."""
    return prs.slides[n - 1]


def find(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError("no shape named %r on this slide (have: %s)"
                   % (name, ", ".join(sorted(s.name for s in slide.shapes))[:400]))


def has(slide, name):
    return any(sh.name == name for sh in slide.shapes)


def _set_para(p, parts):
    """parts: str, or list of (text, bold) tuples for mixed-weight runs."""
    if isinstance(parts, str):
        parts = [(parts, None)]
    runs = list(p.runs)
    if not runs:
        p.add_run()
        runs = list(p.runs)
    # Reuse run 0 as the style template for any extra runs we need.
    base = runs[0]
    base.text = parts[0][0]
    if parts[0][1] is not None:
        base.font.bold = parts[0][1]
    keep = base
    for text, bold in parts[1:]:
        newr = copy.deepcopy(base._r)
        keep._r.addnext(newr)
        keep = p.runs[list(p.runs).index(keep) + 1]
        keep.text = text
        if bold is not None:
            keep.font.bold = bold
    # drop leftover original runs beyond what we wrote
    wanted = len(parts)
    for r in list(p.runs)[wanted:]:
        r._r.getparent().remove(r._r)


def txt(shape, text):
    """Single paragraph. Keeps run 0's formatting."""
    tf = shape.text_frame
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    _set_para(tf.paragraphs[0], text)


def lst(shape, items):
    """Multi-paragraph. Clones paragraph 0 so bullets/spacing survive.

    Each item is a str or a list of (text, bold) tuples.
    """
    if not items:
        return
    tf = shape.text_frame
    template = copy.deepcopy(tf.paragraphs[0]._p)
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    prev = tf.paragraphs[0]._p
    for _ in items[1:]:
        newp = copy.deepcopy(template)
        prev.addnext(newp)
        prev = newp
    for p, item in zip(tf.paragraphs, items):
        _set_para(p, item)


def kill(slide, *names):
    """Remove shapes by name. Silently skips names that aren't present."""
    gone = []
    for name in names:
        for sh in list(slide.shapes):
            if sh.name == name:
                sh._element.getparent().remove(sh._element)
                gone.append(name)
    return gone


def fill(shape, hexcolor):
    from pptx.dml.color import RGBColor
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(hexcolor)


def fontcolor(shape, hexcolor):
    from pptx.dml.color import RGBColor
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor.from_string(hexcolor)


_NOTES_SP = (
    '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<p:nvSpPr><p:cNvPr id="2" name="Notes Placeholder 1"/>'
    '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
    '<p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr>'
    '<p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>'
)


def notes(slide, text):
    """Set speaker notes, creating the body placeholder if the notes master
    doesn't supply one (this deck's doesn't, so notes_text_frame is None)."""
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    if tf is None:
        from pptx.oxml import parse_xml
        ns.shapes._spTree.append(parse_xml(_NOTES_SP))
        tf = ns.notes_text_frame
    tf.text = text


def dump(slide, label=""):
    E = 914400
    print("--- %s (%d shapes)" % (label, len(list(slide.shapes))))
    for sh in slide.shapes:
        t = sh.text_frame.text.strip().replace("\n", " / ")[:40] if sh.has_text_frame else ""
        print("   %-18s L%6.2f T%5.2f W%5.2f H%4.2f | %s"
              % ((sh.name or "")[:18], sh.left / E, sh.top / E,
                 sh.width / E, sh.height / E, t))
